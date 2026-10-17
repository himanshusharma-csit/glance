#!/usr/bin/env python3
"""
GLANCE Analytics Dashboard — export_results.py
Author: Dr. Himanshu Sharma

Turns the GLANCE results workbook into the small JSON file the dashboard
reads:

    python export_results.py --xlsx GLANCE_Results.xlsx --out data

Sheets it understands
---------------------
WordThresholdAblation   (required)  Dataset | Word Threshold | Attention | Epoch |
                                    Batch | BLEU-1..4 | METEOR | ROUGE | CIDEr
                                    Merged cells (Dataset / Threshold / Attention)
                                    are forward-filled. A missing batch size is
                                    inferred from its position (32 → 64 → 128)
                                    and reported under "notes".
AttentionAblation       (optional)  Exported as generic rows; shown once any
                                    metric cell is filled in.
Baselines               (optional)  Dataset | Method | BLEU-1 | BLEU-2 | BLEU-3 |
                                    BLEU-4 | METEOR | ROUGE | CIDEr  (+ optional
                                    Source). Same scale as the main sheet.

A metric that is 0 (or empty) in every run is marked "invalid" and kept out of
charts — e.g. ROUGE-L, which currently reads 0.00 everywhere.

Requires: Python 3.8+, openpyxl.
"""
import argparse
import json
import math
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  pip install openpyxl")

METRICS = ["BLEU-1", "BLEU-2", "BLEU-3", "BLEU-4", "METEOR", "ROUGE", "CIDEr"]
BATCH_CYCLE = [32, 64, 128]

# How the workbook names things -> dashboard keys
DATASETS = {
    "pascal": {"match": r"pascal", "name": "PASCAL 1K-Hindi", "short": "PASCAL", "images": "1,000 images"},
    "flickr": {"match": r"flickr", "name": "Flickr 8K-Hindi", "short": "Flickr", "images": "8,000 images"},
}
ATTENTION = {
    "local": {"match": r"local", "model": "GLANCE-W", "label": "Local (word query)"},
    "global": {"match": r"global", "model": "GLANCE-WH", "label": "Global (word + hidden query)"},
}


def norm(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = re.sub(r"\s+", " ", v).strip()
        return v or None
    return v


def num(v):
    if v is None or isinstance(v, bool):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(f) or math.isinf(f) else round(f, 6)


def key_for(text, table):
    t = (text or "").lower()
    for k, spec in table.items():
        if re.search(spec["match"], t):
            return k
    return None


def sheet_rows(ws):
    """Rows as lists with merged cells filled from their top-left value."""
    grid = [[c.value for c in row] for row in ws.iter_rows()]
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[r - 1][c - 1] = v
    return grid


def find_header(grid, must=("BLEU-1",)):
    for i, row in enumerate(grid):
        cells = [str(norm(c) or "") for c in row]
        if all(m in cells for m in must):
            return i, {c: j for j, c in enumerate(cells) if c}
    return None, {}


def read_threshold_sheet(ws, notes):
    grid = sheet_rows(ws)
    hi, col = find_header(grid, ("Word Threshold", "BLEU-1"))
    if hi is None:
        sys.exit(f"Sheet '{ws.title}': header row with 'Word Threshold' and 'BLEU-1' not found")
    runs, prev = [], {}
    for r, row in enumerate(grid[hi + 1:], start=hi + 2):
        get = lambda name: norm(row[col[name]]) if name in col and col[name] < len(row) else None
        vals = {m: num(get(m)) for m in METRICS}
        if all(v is None for v in vals.values()):
            continue  # spacer row
        ds = key_for(str(get("Dataset") or prev.get("ds_raw") or ""), DATASETS)
        att = key_for(str(get("Attention") or prev.get("att_raw") or ""), ATTENTION)
        thr = num(get("Word Threshold"))
        batch = num(get("Batch"))
        if ds is None or att is None or thr is None:
            notes.append(f"Row {r} skipped: could not read dataset/attention/threshold")
            continue
        if batch is None:
            # position inside its (dataset, attention, threshold) block
            k = sum(1 for x in runs if x["dataset"] == ds and x["attention"] == att and x["threshold"] == int(thr))
            batch = BATCH_CYCLE[k] if k < len(BATCH_CYCLE) else None
            notes.append(f"Row {r} ({DATASETS[ds]['short']}, {ATTENTION[att]['model']}, threshold {int(thr)}): "
                         f"batch size empty in the workbook — inferred as {batch} from row order")
        runs.append({
            "dataset": ds, "attention": att, "threshold": int(thr),
            "batch": int(batch) if batch is not None else None,
            "epoch": int(num(get("Epoch"))) if num(get("Epoch")) is not None else None,
            "metrics": vals, "row": r,
        })
        prev = {"ds_raw": get("Dataset") or prev.get("ds_raw"), "att_raw": get("Attention") or prev.get("att_raw")}
    return runs


def read_generic_sheet(ws):
    grid = sheet_rows(ws)
    hi, col = find_header(grid, ("BLEU-1",))
    if hi is None:
        return {"columns": [], "rows": [], "filled": False}
    cols = [c for c, _ in sorted(col.items(), key=lambda x: x[1])]
    rows = []
    for row in grid[hi + 1:]:
        rec = {c: norm(row[j]) if j < len(row) else None for c, j in col.items()}
        if all(v is None for v in rec.values()):
            continue
        for m in METRICS:
            if m in rec:
                rec[m] = num(rec[m])
        rows.append(rec)
    filled = any(num(r.get(m)) is not None for r in rows for m in METRICS if m in r)
    return {"columns": cols, "rows": rows, "filled": filled}


def read_baselines(ws):
    grid = sheet_rows(ws)
    hi, col = find_header(grid, ("Method", "BLEU-1"))
    if hi is None:
        return []
    out = []
    for row in grid[hi + 1:]:
        get = lambda name: norm(row[col[name]]) if name in col and col[name] < len(row) else None
        ds = key_for(str(get("Dataset") or ""), DATASETS)
        if not ds or not get("Method"):
            continue
        out.append({"dataset": ds, "method": str(get("Method")), "source": get("Source"),
                    "metrics": {m: num(get(m)) for m in METRICS}})
    return out


def main():
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=str(here / "GLANCE_Results.xlsx"), help="results workbook")
    ap.add_argument("--out", default=str(here / "data"), help="output folder (default: ./data)")
    a = ap.parse_args()

    src = Path(a.xlsx)
    if not src.exists():
        sys.exit(f"Workbook not found: {src}")
    wb = openpyxl.load_workbook(src, data_only=True)
    notes = []

    ws_name = next((n for n in wb.sheetnames if n.lower().startswith("wordthreshold")), None)
    if not ws_name:
        sys.exit("Sheet 'WordThresholdAblation' not found")
    runs = read_threshold_sheet(wb[ws_name], notes)

    invalid = [m for m in METRICS if all(not r["metrics"].get(m) for r in runs)]
    for m in invalid:
        notes.append(f"{m} is 0 or empty in every run — treated as not computed and kept out of charts")

    att_sheet = next((n for n in wb.sheetnames if n.lower().startswith("attentionablation")), None)
    attention_ablation = read_generic_sheet(wb[att_sheet]) if att_sheet else {"columns": [], "rows": [], "filled": False}
    base_sheet = next((n for n in wb.sheetnames if n.lower().startswith("baseline")), None)
    baselines = read_baselines(wb[base_sheet]) if base_sheet else []

    datasets = []
    for k, spec in DATASETS.items():
        rs = [r for r in runs if r["dataset"] == k]
        if not rs:
            continue
        datasets.append({
            "key": k, "name": spec["name"], "short": spec["short"], "images": spec["images"],
            "n_runs": len(rs),
            "thresholds": sorted({r["threshold"] for r in rs}),
            "batches": sorted({r["batch"] for r in rs if r["batch"] is not None}),
            "availability": {
                "threshold_ablation": True,
                "both_attentions": len({r["attention"] for r in rs}) == 2,
                "attention_ablation": attention_ablation["filled"],
                "baselines": any(b["dataset"] == k for b in baselines),
            },
        })

    out = {
        "project": "GLANCE",
        "title": "Global-Local Attention Network for Cross-modal Alignment in Hindi Image Captioning",
        "exported": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "source_file": src.name,
        "metrics": METRICS,
        "invalid_metrics": invalid,
        "attention": {k: {"model": v["model"], "label": v["label"]} for k, v in ATTENTION.items()},
        "datasets": datasets,
        "runs": runs,
        "attention_ablation": attention_ablation,
        "baselines": baselines,
        "notes": notes,
    }
    dst = Path(a.out)
    dst.mkdir(parents=True, exist_ok=True)
    p = dst / "results.json"
    p.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(runs)} runs from {src.name} → {p}  ({p.stat().st_size / 1024:.0f} KB)")
    for n in notes:
        print("  note:", n)


if __name__ == "__main__":
    main()
